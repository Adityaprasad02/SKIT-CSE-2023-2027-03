from unittest.mock import MagicMock, patch

import pytest

from docx.opc.exceptions import PackageNotFoundError

from resume.parsing.docx_parser import (
    extract_text_from_docx,
    extract_lines_from_docx,
)


# =========================================
# extract_text_from_docx() Tests
# =========================================

@patch("resume.parsing.docx_parser.Document")
def test_extract_text_from_docx_paragraphs(mock_document):

    paragraph1 = MagicMock()
    paragraph1.text = "Python Developer"

    paragraph2 = MagicMock()
    paragraph2.text = "Machine Learning Engineer"

    document = MagicMock()
    document.paragraphs = [paragraph1, paragraph2]
    document.tables = []

    mock_document.return_value = document

    result = extract_text_from_docx("resume.docx")

    assert result == (
        "Python Developer\n"
        "Machine Learning Engineer"
    )


@patch("resume.parsing.docx_parser.Document")
def test_extract_text_from_docx_ignores_empty_paragraphs(mock_document):

    paragraph1 = MagicMock()
    paragraph1.text = "Python"

    paragraph2 = MagicMock()
    paragraph2.text = "   "

    document = MagicMock()
    document.paragraphs = [paragraph1, paragraph2]
    document.tables = []

    mock_document.return_value = document

    result = extract_text_from_docx("resume.docx")

    assert result == "Python"


@patch("resume.parsing.docx_parser.Document")
def test_extract_text_from_docx_extracts_tables(mock_document):

    cell1 = MagicMock()
    cell1.text = "Python"

    cell2 = MagicMock()
    cell2.text = "Machine Learning"

    row = MagicMock()
    row.cells = [cell1, cell2]

    table = MagicMock()
    table.rows = [row]

    document = MagicMock()
    document.paragraphs = []
    document.tables = [table]

    mock_document.return_value = document

    result = extract_text_from_docx("resume.docx")

    assert result == "Python | Machine Learning"


@patch("resume.parsing.docx_parser.Document")
def test_extract_text_from_docx_no_readable_text(mock_document):

    paragraph = MagicMock()
    paragraph.text = "   "

    document = MagicMock()
    document.paragraphs = [paragraph]
    document.tables = []

    mock_document.return_value = document

    with pytest.raises(ValueError, match="No readable text found"):
        extract_text_from_docx("empty.docx")


@patch("resume.parsing.docx_parser.Document")
def test_extract_text_from_docx_invalid_file(mock_document):

    mock_document.side_effect = PackageNotFoundError("Invalid DOCX")

    with pytest.raises(RuntimeError, match="Not a valid DOCX file"):
        extract_text_from_docx("broken.docx")


# =========================================
# extract_lines_from_docx() Tests
# =========================================

@patch("resume.parsing.docx_parser.Document")
def test_extract_lines_from_docx(mock_document):

    run = MagicMock()
    run.text = "Python Developer"
    run.bold = True
    run.italic = False
    run.font.size.pt = 14.0

    paragraph = MagicMock()
    paragraph.text = "Python Developer"
    paragraph.runs = [run]
    paragraph.style.name = "Normal"

    document = MagicMock()
    document.paragraphs = [paragraph]
    document.tables = []

    mock_document.return_value = document

    result = extract_lines_from_docx("resume.docx")

    assert len(result) == 1

    line = result[0]

    assert line.text == "Python Developer"
    assert line.page_number == 1
    assert line.font_size == 14.0
    assert line.is_bold is True
    assert line.is_italic is False
    assert line.line_index == 0


@patch("resume.parsing.docx_parser.Document")
def test_extract_lines_from_docx_heading_style(mock_document):

    run = MagicMock()
    run.text = "Education"
    run.bold = False
    run.italic = False
    run.font.size = None

    paragraph = MagicMock()
    paragraph.text = "Education"
    paragraph.runs = [run]
    paragraph.style.name = "Heading 1"

    document = MagicMock()
    document.paragraphs = [paragraph]
    document.tables = []

    mock_document.return_value = document

    result = extract_lines_from_docx("resume.docx")

    line = result[0]

    assert line.is_bold is True
    assert line.font_size == 14.0


@patch("resume.parsing.docx_parser.Document")
def test_extract_lines_from_docx_title_style(mock_document):

    run = MagicMock()
    run.text = "Aditya Singh"
    run.bold = False
    run.italic = False
    run.font.size = None

    paragraph = MagicMock()
    paragraph.text = "Aditya Singh"
    paragraph.runs = [run]
    paragraph.style.name = "Title"

    document = MagicMock()
    document.paragraphs = [paragraph]
    document.tables = []

    mock_document.return_value = document

    result = extract_lines_from_docx("resume.docx")

    assert result[0].is_bold is True
    assert result[0].font_size == 14.0


@patch("resume.parsing.docx_parser.Document")
def test_extract_lines_from_docx_extracts_table(mock_document):

    cell1 = MagicMock()
    cell1.text = "Python"

    cell2 = MagicMock()
    cell2.text = "SQL"

    row = MagicMock()
    row.cells = [cell1, cell2]

    table = MagicMock()
    table.rows = [row]

    document = MagicMock()
    document.paragraphs = []
    document.tables = [table]

    mock_document.return_value = document

    result = extract_lines_from_docx("resume.docx")

    assert len(result) == 1

    line = result[0]

    assert line.text == "Python | SQL"
    assert line.page_number == 1
    assert line.font_size == 0.0
    assert line.is_bold is False
    assert line.is_italic is False
    assert line.line_index == 0


@patch("resume.parsing.docx_parser.Document")
def test_extract_lines_from_docx_line_indexes(mock_document):

    run1 = MagicMock()
    run1.text = "Summary"
    run1.bold = False
    run1.italic = False
    run1.font.size = None

    run2 = MagicMock()
    run2.text = "Skills"
    run2.bold = False
    run2.italic = False
    run2.font.size = None

    paragraph1 = MagicMock()
    paragraph1.text = "Summary"
    paragraph1.runs = [run1]
    paragraph1.style.name = "Normal"

    paragraph2 = MagicMock()
    paragraph2.text = "Skills"
    paragraph2.runs = [run2]
    paragraph2.style.name = "Normal"

    document = MagicMock()
    document.paragraphs = [paragraph1, paragraph2]
    document.tables = []

    mock_document.return_value = document

    result = extract_lines_from_docx("resume.docx")

    assert result[0].line_index == 0
    assert result[1].line_index == 1


@patch("resume.parsing.docx_parser.Document")
def test_extract_lines_from_docx_no_readable_lines(mock_document):

    paragraph = MagicMock()
    paragraph.text = "   "

    document = MagicMock()
    document.paragraphs = [paragraph]
    document.tables = []

    mock_document.return_value = document

    with pytest.raises(ValueError, match="No readable text found"):
        extract_lines_from_docx("empty.docx")


@patch("resume.parsing.docx_parser.Document")
def test_extract_lines_from_docx_invalid_file(mock_document):

    mock_document.side_effect = PackageNotFoundError("Invalid DOCX")

    with pytest.raises(RuntimeError, match="Not a valid DOCX file"):
        extract_lines_from_docx("broken.docx")