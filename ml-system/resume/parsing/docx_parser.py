from docx import Document
from docx.opc.exceptions import PackageNotFoundError

from resume.parsing.models import RawLine


def extract_text_from_docx(file_path: str) -> str:
    """Plain-text extraction. Kept for backward compatibility."""

    try:
        document = Document(file_path)
    except PackageNotFoundError as error:
        raise RuntimeError(f"Not a valid DOCX file: {error}") from error

    text_parts = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            text_parts.append(text)

    for table in document.tables:
        for row in table.rows:
            row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_text:
                text_parts.append(" | ".join(row_text))

    extracted_text = "\n".join(text_parts)

    if not extracted_text.strip():
        raise ValueError("No readable text found in the DOCX file.")

    return extracted_text.strip()


def extract_lines_from_docx(file_path: str) -> list[RawLine]:
    """Structured extraction mirroring extract_lines_from_pdf's output shape.

    Layout signals here come from two structural sources, neither of which
    is a keyword list:
      - the paragraph's Word style (e.g. "Heading 1", "Title") when the
        author used real heading styles
      - run-level bold/italic/size when they used manual formatting instead
        (very common in resumes built from templates)
    """

    try:
        document = Document(file_path)
    except PackageNotFoundError as error:
        raise RuntimeError(f"Not a valid DOCX file: {error}") from error

    lines: list[RawLine] = []
    line_index = 0

    # DOCX has no "page number" concept without rendering; we use a
    # constant since downstream code only needs page_number for pagination
    # display, not for section logic.
    page_number = 1

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue

        runs = [r for r in paragraph.runs if r.text.strip()]

        if runs:
            first_run = runs[0]
            is_bold = bool(first_run.bold)
            is_italic = bool(first_run.italic)
            font_size = first_run.font.size.pt if first_run.font.size else 0.0
        else:
            is_bold = False
            is_italic = False
            font_size = 0.0

        style_name = (paragraph.style.name or "").lower() if paragraph.style else ""

        # A named heading style is a strong structural signal in its own
        # right — bump the effective "size" so header_detector treats it
        # as bold/large even if no manual run formatting was applied.
        if "heading" in style_name or "title" in style_name:
            is_bold = True
            font_size = max(font_size, 14.0)

        lines.append(
            RawLine(
                text=text,
                page_number=page_number,
                font_size=round(font_size, 1),
                is_bold=is_bold,
                is_italic=is_italic,
                line_index=line_index,
            )
        )
        line_index += 1

    for table in document.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                lines.append(
                    RawLine(
                        text=row_text,
                        page_number=page_number,
                        font_size=0.0,
                        is_bold=False,
                        is_italic=False,
                        line_index=line_index,
                    )
                )
                line_index += 1

    if not lines:
        raise ValueError("No readable text found in the DOCX file.")

    return lines